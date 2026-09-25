"""Lembar validasi manual Titik Aman Semu v6 (TAS, aturan utama) untuk keempat level.

    python -m scripts.lembar_validasi_TAS

Keluaran (output_bab4/validasi/):
  validasi_TAS_v6.xlsx   lembar "Validasi TAS" (satu baris per grid TAS per level), "Hotspot" (DBSCAN 350 m pada
                         grid unik yang perlu divalidasi/dikonfirmasi), dan "Petunjuk"
  peta/<level>_<id_grid>.png   peta per TAS: grid, TES terdekat, garis lurus, rute jaringan, ruas snapping,
                         ruas tertutup; untuk grid "dipicu"/"diperparah banjir" juga rute Baseline ke TES yang sama

Status TAS, T_ideal, T_aktual, DI_t, dan atribusi banjir diambil dari hasil TERKUNCI (data/locked/). Rute
dihitung ulang dengan aturan yang sama (graf jalan level tersebut, snapping ke ruas ≤ 300 m, batas minimum
50 m); T_aktual hasil hitung ulang dicocokkan dengan nilai terkunci (kolom "Cek T_aktual").

Isian awal otomatis (aturan ditetapkan peneliti sebelum hasil v6 terlihat) dibuat oleh `prefill`; teks v5
dibaca dari validasi_TAS_v5_terisi.xlsx. Setelah diisi, lembar direkap dengan
`python -m scripts.rekap_validasi_TAS`.
"""
import json
import logging
import sys
from pathlib import Path

import numpy as np
import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.rekap_validasi_TAS import (KODE, KODE_PILIHAN, LEMBAR, SHEET, SHEET_HOTSPOT, STATUS_ISIAN,
                                        VALID_DARI_KODE, VALIDASI_DIR)

LOCKED = PROJECT_ROOT / "data" / "locked"
LEMBAR_V5 = VALIDASI_DIR / "validasi_TAS_v5_terisi.xlsx"
NAME_COLS = ["Nama_Objek", "nama", "Nama", "NAMA", "Name", "NAME", "REMARK", "Fasilitas", "KETERANGAN", "NAMA_UNSUR"]
KAT_LABEL = {"pendidikan": "Pendidikan", "kesehatan": "Kesehatan", "pemerintahan": "Pemerintahan",
             "ibadah": "Tempat Ibadah", "gor": "GOR/Gedung Serbaguna"}
HOTSPOT_EPS_M = 350.0
TOL_MENIT = 1.0                                     # sama dengan thesis.ATRIBUSI_TOL_MENIT
TEKS_UMUM = "opsi jalan yang disediakan oleh osm terbatas"   # kalimat umum v5 (Sedang/Tinggi)
OTO, KONF, VAL = STATUS_ISIAN                       # otomatis / perlu konfirmasi peneliti / perlu divalidasi

# Aturan kata kunci teks v5 → kode (diperiksa berurutan I, F, A, B, C, D, H, E)
KATA_KUNCI = [
    ("I", lambda t: "pasti ada kesalahan" in t),
    ("F", lambda t: "snapping" in t),
    ("A", lambda t: ("sungai" in t or "waduk" in t) and "jembatan" in t),
    ("B", lambda t: "bandara" in t and "memang" in t),
    ("C", lambda t: any(k in t for k in ("perbukitan", "lembah", "kontur"))),
    ("D", lambda t: "pesisir" in t),
    ("H", lambda t: "terscraping" in t),
    ("E", lambda t: "osm" in t or "google" in t),
]
KOLOM = ["id_grid", "id_grid_asli", "level", "Hotspot", "Lat centroid", "Lon centroid", "TES terdekat", "Kategori TES",
         "Lat TES", "Lon TES", "T_ideal (menit)", "T_aktual (menit)", "DI_t", "Cek T_aktual (hitung ulang)",
         "Atribusi banjir", "Waktu Baseline ke TES sama (menit)", "Tambahan waktu akibat penutupan (menit)",
         "Google Maps centroid", "Google Maps TES", "Peta", "Kode utama", "Kode tambahan", "Valid", "Status isian",
         "Penyebab (teks)", "Catatan"]
KOLOM_ISI = ("Kode utama", "Kode tambahan", "Status isian", "Penyebab (teks)", "Catatan")


def _teks(v) -> str:
    return "" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v).strip()


def kode_dari_teks(teks: str) -> list:
    """Kode yang cocok menurut aturan kata kunci, berurutan I, F, A, B, C, D, H, E."""
    t = _teks(teks).lower()
    return [k for k, f in KATA_KUNCI if f(t)] if t else []


def pilihan(kode: str) -> str:
    """Huruf kode → teks dropdown ("A" → "A Sungai/waduk")."""
    return f"{kode} {KODE[kode]}" if kode in KODE else ""


def prefill(rows: pd.DataFrame, v5: pd.DataFrame) -> pd.DataFrame:
    """Isian awal otomatis kolom Kode utama, Kode tambahan, Status isian, Penyebab (teks), dan Catatan.

    rows: baris lembar v6 (id_grid, level, Lat TES, Lon TES, T_aktual (menit), Atribusi banjir).
    v5  : lembar v5 terisi (id_grid, level, Lat TES, Lon TES, T_aktual (menit), Penyebab, Catatan).
    Urutan prioritas (a)–(e) mengikuti instruksi Putaran 6 (lihat docs/CATATAN_TEMUAN.md P6-4)."""
    key5 = {(int(r["id_grid"]), r["level"]): r for _, r in v5.iterrows()}
    out = pd.DataFrame("", index=rows.index, columns=list(KOLOM_ISI), dtype=object)

    def arsip(r5, alasan):
        isi = "; ".join(f"{k}: {_teks(r5[k])}" for k in ("Penyebab", "Catatan") if _teks(r5.get(k)))
        return f"Arsip v5 ({alasan}; tidak dipakai untuk kode) — {isi}" if isi else ""

    def v5_cocok(r):
        """(baris v5, TES sama?, selisih T_aktual v6 − v5)."""
        r5 = key5.get((int(r["id_grid"]), r["level"]))
        if r5 is None:
            return None, False, np.nan
        same = abs(r5["Lat TES"] - r["Lat TES"]) < 1e-6 and abs(r5["Lon TES"] - r["Lon TES"]) < 1e-6
        return r5, same, float(r["T_aktual (menit)"] - r5["T_aktual (menit)"])

    def put(i, o):
        for k, v in o.items():
            out.at[i, k] = v

    base_code = {}
    for i, r in rows[rows["level"] == "Baseline"].iterrows():          # (a) dan (e) untuk Baseline
        r5, same, d = v5_cocok(r)
        o = {"Status isian": VAL}
        if r5 is not None and same and abs(d) <= TOL_MENIT:
            o["Penyebab (teks)"], o["Catatan"] = _teks(r5.get("Penyebab")), _teks(r5.get("Catatan"))
            ks = kode_dari_teks(f"{o['Penyebab (teks)']} {o['Catatan']}")
            if "F" in ks:
                o["Catatan"] = ("[v6] Kode F dari teks v5 tidak dibawa: snapping sudah diperbaiki, tetapi grid masih "
                                "TAS. " + o["Catatan"]).strip()
            elif ks:
                o.update({"Kode utama": pilihan(ks[0]), "Kode tambahan": pilihan(ks[1]) if len(ks) > 1 else "",
                          "Status isian": KONF})
        elif r5 is not None:
            o["Catatan"] = arsip(r5, "TES terdekat berbeda" if not same else f"T_aktual berubah {d:+.2f} menit")
        put(i, o)
        base_code[int(r["id_grid"])] = (out.at[i, "Kode utama"], out.at[i, "Kode tambahan"])

    for i, r in rows[rows["level"] != "Baseline"].iterrows():          # (b)–(e) untuk Rendah/Sedang/Tinggi
        atr = _teks(r.get("Atribusi banjir"))
        bu, bt = base_code.get(int(r["id_grid"]), ("", ""))
        o = {"Status isian": VAL}
        if atr == "dipicu banjir":
            o.update({"Kode utama": pilihan("G"), "Status isian": OTO})
        elif atr == "diperparah banjir":
            o.update({"Kode utama": bu, "Kode tambahan": pilihan("G"), "Status isian": OTO if bu else VAL})
        elif atr == "tidak berubah":
            o.update({"Kode utama": bu, "Kode tambahan": bt if bt != pilihan("G") else "",
                      "Status isian": OTO if bu else VAL})
        r5, same, d = v5_cocok(r)
        if r5 is not None:
            if TEKS_UMUM in _teks(r5.get("Penyebab")).lower():
                o["Catatan"] = arsip(r5, "kalimat umum v5")
            elif same and abs(d) <= TOL_MENIT:
                o["Penyebab (teks)"], o["Catatan"] = _teks(r5.get("Penyebab")), _teks(r5.get("Catatan"))
            else:
                o["Catatan"] = arsip(r5, "TES terdekat berbeda" if not same else f"T_aktual berubah {d:+.2f} menit")
        put(i, o)
    return out


def hotspots(xy: np.ndarray, eps: float = HOTSPOT_EPS_M) -> np.ndarray:
    """Label DBSCAN (eps meter, min_samples = 1: setiap grid masuk satu hotspot, termasuk grid tunggal)."""
    from sklearn.cluster import DBSCAN
    if len(xy) == 0:
        return np.zeros(0, int)
    return DBSCAN(eps=eps, min_samples=1).fit_predict(np.asarray(xy, float))


def tes_name(row) -> str:
    for c in NAME_COLS:
        if c in row.index and pd.notna(row[c]) and str(row[c]).strip():
            return str(row[c]).strip()
    return "(tanpa nama)"


def gmaps(lat, lon) -> str:
    return f"https://www.google.com/maps/search/?api=1&query={lat:.6f},{lon:.6f}"


def gmaps_satelit(lat, lon) -> str:
    return f"https://www.google.com/maps/@{lat:.6f},{lon:.6f},18z/data=!3m1!1e3"


def _map(png, grid_geom, tes_pt, route, gxy, roads_lv, closed_lv, route_base=None):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    pts = [gxy, tes_pt] + (route["xy"] if route else []) + (route_base["xy"] if route_base else [])
    xs, ys = [p[0] for p in pts], [p[1] for p in pts]
    pad = 150.0
    x0, x1, y0, y1 = min(xs) - pad, max(xs) + pad, min(ys) - pad, max(ys) + pad
    fig, ax = plt.subplots(figsize=(4.2, 4.2), dpi=110)
    sub = roads_lv.cx[x0:x1, y0:y1]
    sc = closed_lv[sub.index.values]
    if (~sc).any():
        sub[~sc].plot(ax=ax, color="#9ca3af", linewidth=0.6)
    if sc.any():
        sub[sc].plot(ax=ax, color="#dc2626", linewidth=0.8, linestyle="--")
    for poly in getattr(grid_geom, "geoms", [grid_geom]):
        ax.fill(*poly.exterior.xy, facecolor="#a855f7", alpha=0.35, edgecolor="#6b21a8", linewidth=1)
    ax.plot([gxy[0], tes_pt[0]], [gxy[1], tes_pt[1]], color="#111827", linestyle=":", linewidth=1.2)
    if route_base:
        net = route_base["xy"][1:-1]
        ax.plot([p[0] for p in net], [p[1] for p in net], color="#0d9488", linewidth=2.6, alpha=0.55)
    if route:
        p = route["xy"]
        ax.plot([q[0] for q in p[1:-1]], [q[1] for q in p[1:-1]], color="#2563eb", linewidth=1.6)
        for a_, b_ in ((p[0], p[1]), (p[-2], p[-1])):                     # ruas snapping
            ax.plot([a_[0], b_[0]], [a_[1], b_[1]], color="#f97316", linewidth=1.4, linestyle="--")
    ax.plot(*tes_pt, marker="*", markersize=12, color="#16a34a")
    ax.set_xlim(x0, x1)
    ax.set_ylim(y0, y1)
    ax.set_aspect("equal")
    ax.set_xticks([])
    ax.set_yticks([])
    ax.set_title("ungu = grid TAS; bintang = TES terdekat; ··· garis lurus; biru = rute jaringan\n"
                 "oranye putus = ruas snapping; merah putus = ruas ditutup"
                 + ("; hijau tebal = rute Baseline" if route_base else ""), fontsize=6.3)
    fig.tight_layout()
    fig.savefig(png)
    plt.close(fig)


def build(out_dir: Path = VALIDASI_DIR, lembar_v5: Path = LEMBAR_V5):
    from pyproj import Transformer
    from backend import thesis as T
    from backend.engine import Cfg, load_data, load_road_network, road_closed_mask
    logging.basicConfig(level=logging.WARNING)
    R = json.loads((LOCKED / "thesis_results.json").read_text(encoding="utf-8"))
    grid = pd.read_csv(LOCKED / "thesis_grid_results.csv.gz")
    t_pen = float(R["t_pen"])
    cfg = Cfg(base_dir=str(PROJECT_ROOT / "data"))
    gdf = load_data(cfg)
    roads, tes = load_road_network(cfg)
    to_wgs = Transformer.from_crs(cfg.target_crs, "EPSG:4326", always_xy=True)
    speed = cfg.walking_speed_m_per_min
    peta_dir = out_dir / "peta"
    peta_dir.mkdir(parents=True, exist_ok=True)
    for old in peta_dir.glob("*.png"):
        old.unlink()
    gxy_all = np.column_stack([gdf["cx"].values, gdf["cy"].values])
    acc_b = T.compute_level_accessibility(gdf, roads, tes, 0.0, cfg, t_pen=t_pen)
    graph_b = acc_b["graph"]
    rows = []
    for lv in T.LEVELS:
        key = lv["key"]
        sel = np.where(grid[f"tas_status_{key}"].values == 1)[0]
        if not len(sel):
            continue
        acc = acc_b if key == "baseline" else T.compute_level_accessibility(gdf, roads, tes, lv["intensity"], cfg,
                                                                              t_pen=t_pen)
        tes_v = acc["tes_v"].set_index("id_tes")
        closed = road_closed_mask(roads, T.SKENARIO, lv["intensity"], cfg)
        ids = grid[f"id_tes_terdekat_{key}"].values[sel]
        txy = np.column_stack([tes_v.loc[ids].geometry.x.values, tes_v.loc[ids].geometry.y.values])
        routes = T.tas_routes(acc["graph"], gxy_all[sel], txy)
        atr_col = f"atribusi_banjir_{key}"
        atr = grid[atr_col].values[sel] if atr_col in grid.columns else np.full(len(sel), None)
        need_b = np.array([a in ("dipicu banjir", "diperparah banjir") for a in atr])
        routes_b = [None] * len(sel)
        if need_b.any():
            rb = T.tas_routes(graph_b, gxy_all[sel][need_b], txy[need_b])
            for j, r in zip(np.where(need_b)[0], rb):
                routes_b[j] = r
        for j, gi in enumerate(sel):
            gxy = gxy_all[gi]
            r = routes[j]
            t_akt_ulang = min(max(r["jarak_m"], T.TAS_MIN_EUCLID_M) / speed, t_pen) if r else t_pen
            t_akt = float(grid.at[gi, f"t_aktual_{key}"])
            lon_g, lat_g = to_wgs.transform(*gxy)
            lon_t, lat_t = to_wgs.transform(*txy[j])
            png = peta_dir / f"{key}_{int(grid.at[gi, 'id_grid'])}.png"
            _map(png, gdf.geometry.iloc[gi], tuple(txy[j]), r, tuple(gxy), roads, closed, routes_b[j])
            trow = tes_v.loc[ids[j]]
            num = lambda c: (float(grid.at[gi, c]) if c in grid.columns and pd.notna(grid.at[gi, c]) else None)
            rows.append({
                "id_grid": int(grid.at[gi, "id_grid"]), "id_grid_asli": str(grid.at[gi, "id_grid_asli"]),
                "level": lv["label"], "Hotspot": None, "Lat centroid": round(lat_g, 6), "Lon centroid": round(lon_g, 6),
                "TES terdekat": tes_name(trow), "Kategori TES": KAT_LABEL.get(trow["kategori"], trow["kategori"]),
                "Lat TES": round(lat_t, 6), "Lon TES": round(lon_t, 6),
                "T_ideal (menit)": float(grid.at[gi, f"t_ideal_{key}"]), "T_aktual (menit)": t_akt,
                "DI_t": float(grid.at[gi, f"di_{key}"]),
                "Cek T_aktual (hitung ulang)": "sama" if abs(t_akt_ulang - t_akt) < 1e-2 else f"beda ({t_akt_ulang:.2f})",
                "Atribusi banjir": atr[j] if isinstance(atr[j], str) else None,
                "Waktu Baseline ke TES sama (menit)": num(f"waktu_baseline_tes_sama_{key}"),
                "Tambahan waktu akibat penutupan (menit)": num(f"tambahan_waktu_{key}"),
                "Google Maps centroid": gmaps(lat_g, lon_g), "Google Maps TES": gmaps(lat_t, lon_t),
                "Peta": f"peta/{png.name}",
                "_x": gxy[0], "_y": gxy[1],
            })
    df = pd.DataFrame(rows)
    v5 = pd.read_excel(lembar_v5, sheet_name=SHEET)
    df[list(KOLOM_ISI)] = prefill(df, v5)
    df["Valid"] = None
    df, hs = add_hotspots(df)
    write_xlsx(df[KOLOM], hs, out_dir / LEMBAR.name)
    return df[KOLOM], hs


def add_hotspots(df: pd.DataFrame, eps: float = HOTSPOT_EPS_M):
    """Kolom Hotspot pada lembar utama dan tabel lembar Hotspot. Hotspot = klaster DBSCAN pada grid unik yang
    punya baris "perlu divalidasi" atau "perlu konfirmasi peneliti" (koordinat UTM _x/_y); dinomori H01, H02, …
    menurut jumlah grid (menurun), lalu id_grid terkecil."""
    df = df.copy()
    perlu = df[df["Status isian"].isin([VAL, KONF])].drop_duplicates("id_grid")
    uniq = pd.DataFrame({"id_grid": perlu["id_grid"].values, "lab": hotspots(perlu[["_x", "_y"]].values, eps)})
    order = uniq.groupby("lab")["id_grid"].agg(["size", "min"]).sort_values(["size", "min"], ascending=[False, True])
    name = {lab: f"H{i + 1:02d}" for i, lab in enumerate(order.index)}
    df["Hotspot"] = df["id_grid"].map(dict(zip(uniq["id_grid"], uniq["lab"].map(name))))
    rows = []
    for lab in order.index:
        gids = set(uniq.loc[uniq["lab"] == lab, "id_grid"])
        sub = df[df["id_grid"].isin(gids)]
        g1 = sub.drop_duplicates("id_grid")
        lat, lon = float(g1["Lat centroid"].mean()), float(g1["Lon centroid"].mean())
        rows.append({"Hotspot": name[lab], "Jumlah grid": len(gids),
                     "id_grid": ", ".join(str(g) for g in sorted(gids)),
                     "Level": ", ".join(x for x in ["Baseline", "Rendah", "Sedang", "Tinggi"] if x in set(sub["level"])),
                     "TES": "; ".join(sorted(set(sub["TES terdekat"]))), "DI maksimum": float(sub["DI_t"].max()),
                     "Lat": round(lat, 6), "Lon": round(lon, 6), "Google Maps satelit": gmaps_satelit(lat, lon),
                     "Kode utama": None, "Kode tambahan": None, "Catatan": None})
    cols = ["Hotspot", "Jumlah grid", "id_grid", "Level", "TES", "DI maksimum", "Lat", "Lon", "Google Maps satelit",
            "Kode utama", "Kode tambahan", "Catatan"]
    return df, pd.DataFrame(rows, columns=cols)


def _style_sheet(ws, cols, n, link_cols, fill_cols):
    from openpyxl.styles import Alignment, Font, PatternFill
    head = Font(bold=True)
    fill_isi = PatternFill("solid", fgColor="FFF7CC")
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j)
        cell.font = head
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[cell.column_letter].width = max(12, min(40, len(c) + 4))
    for i in range(2, n + 1):
        for c in link_cols:
            cell = ws.cell(row=i, column=cols.index(c) + 1)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="1D4ED8", underline="single")
        for c in fill_cols:
            ws.cell(row=i, column=cols.index(c) + 1).fill = fill_isi


def _cell(v):
    if v is None or (isinstance(v, float) and np.isnan(v)) or (isinstance(v, str) and v == ""):
        return None
    return v


def write_xlsx(df: pd.DataFrame, hs: pd.DataFrame, path: Path):
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    cols = list(df.columns)
    ws.append(cols)
    for r in df.itertuples(index=False):
        ws.append([_cell(v) for v in r])
    n = len(df) + 1
    col = lambda c: ws.cell(row=1, column=cols.index(c) + 1).column_letter
    ku, v = col("Kode utama"), col("Valid")
    for i in range(2, n + 1):                                            # Valid otomatis dari Kode utama
        k = f"LEFT({ku}{i},1)"
        ws[f"{v}{i}"] = (f'=IF({k}="","",IF(ISNUMBER(SEARCH({k},"ABCDG")),"Y",'
                         f'IF(ISNUMBER(SEARCH({k},"EFH")),"T",IF({k}="I","cek",""))))')
    _style_sheet(ws, cols, n, ("Google Maps centroid", "Google Maps TES", "Peta"),
                 ("Kode utama", "Kode tambahan", "Status isian", "Penyebab (teks)", "Catatan"))
    dv_kode = DataValidation(type="list", formula1='"' + ",".join(KODE_PILIHAN) + '"', allow_blank=True)
    dv_stat = DataValidation(type="list", formula1='"' + ",".join(STATUS_ISIAN) + '"', allow_blank=True)
    ws.add_data_validation(dv_kode)
    ws.add_data_validation(dv_stat)
    if n >= 2:
        for c in ("Kode utama", "Kode tambahan"):
            dv_kode.add(f"{col(c)}2:{col(c)}{n}")
        dv_stat.add(f"{col('Status isian')}2:{col('Status isian')}{n}")
    ws.freeze_panes = "E2"

    wh = wb.create_sheet(SHEET_HOTSPOT)
    hcols = list(hs.columns)
    wh.append(hcols)
    for r in hs.itertuples(index=False):
        wh.append([_cell(x) for x in r])
    nh = len(hs) + 1
    _style_sheet(wh, hcols, nh, ("Google Maps satelit",), ("Kode utama", "Kode tambahan", "Catatan"))
    dv_kode2 = DataValidation(type="list", formula1='"' + ",".join(KODE_PILIHAN) + '"', allow_blank=True)
    wh.add_data_validation(dv_kode2)
    if nh >= 2:
        for c in ("Kode utama", "Kode tambahan"):
            L = wh.cell(row=1, column=hcols.index(c) + 1).column_letter
            dv_kode2.add(f"{L}2:{L}{nh}")
    wh.freeze_panes = "B2"

    info = wb.create_sheet("Petunjuk")
    for line in PETUNJUK:
        info.append([line])
    info.column_dimensions["A"].width = 130
    wb.save(path)


PETUNJUK = [
    "Lembar validasi manual Titik Aman Semu v6 (aturan utama: DI ≥ 2, T_ideal ≤ 5 menit, T_aktual ≥ 30 menit).",
    "Satu baris = satu grid TAS pada satu level. Nilai diambil dari data/locked/ (hasil-skripsi-v6; snapping ke ruas).",
    "",
    "KODE (kolom Kode utama dan Kode tambahan; pilih dari daftar):",
    *[f"  {p}" for p in KODE_PILIHAN],
    "",
    "VALID (otomatis dari Kode utama): " + "; ".join(
        f"{'/'.join(k for k in KODE if VALID_DARI_KODE[k] == v)} → {v}" for v in ("Y", "T", "cek")),
    "  Y = TAS nyata (penghalang fisik atau penutupan banjir); T = artefak data/metode; cek = belum terjelaskan.",
    "",
    "ATRIBUSI BANJIR (level Rendah/Sedang/Tinggi; kosong untuk Baseline). TES sama = TES terdekat (Euclidean) identik "
    "dengan Baseline.",
    "  dipicu banjir     : grid bukan TAS di Baseline, TES sama, dan waktu Baseline ke TES itu < 30 menit.",
    "  diperparah banjir : grid juga TAS di Baseline, TES sama, dan T_aktual naik > 1 menit dibanding Baseline.",
    "  tidak berubah     : grid juga TAS di Baseline, TES sama, dan T_aktual berubah ≤ 1 menit.",
    "  lainnya           : selain ketiganya (mis. TES terdekat berubah karena TES Baseline tidak valid di level itu).",
    "  Waktu Baseline ke TES sama dan Tambahan waktu hanya diisi bila TES sama.",
    "",
    "STATUS ISIAN:",
    "  otomatis                  : kode diisi dari aturan (dipicu → G; diperparah → kode Baseline + G; tidak berubah → "
    "kode Baseline).",
    "  perlu konfirmasi peneliti : kode Baseline dipetakan dari teks v5 dengan aturan kata kunci; periksa kembali.",
    "  perlu divalidasi          : belum ada kode; periksa peta dan Google Maps, lalu isi.",
    "  Teks v5 berkode F (snapping) tidak dibawa bila grid masih TAS di v6, karena snapping sudah diperbaiki.",
    "  Kalimat umum v5 untuk Sedang/Tinggi hanya disimpan di Catatan sebagai arsip.",
    "",
    "HOTSPOT: DBSCAN 350 m (min. 1 grid) pada grid unik berstatus perlu divalidasi / perlu konfirmasi. Periksa "
    "hotspot sekali (tautan satelit), isi kodenya di lembar Hotspot, lalu salin ke baris grid terkait bila sesuai.",
    "",
    "PETA (kolom Peta): ungu = grid; bintang = TES terdekat; titik-titik = garis lurus; biru = rute jaringan; oranye "
    "putus = ruas snapping; merah putus = ruas ditutup; hijau tebal = rute Baseline (dipicu/diperparah).",
    "Setelah diisi: python -m scripts.rekap_validasi_TAS",
]


if __name__ == "__main__":
    d, h = build()
    print(d.groupby("level").size().reindex(["Baseline", "Rendah", "Sedang", "Tinggi"]).fillna(0).astype(int))
    print("status isian:", d["Status isian"].value_counts().to_dict())
    print("cek T_aktual:", d["Cek T_aktual (hitung ulang)"].value_counts().to_dict())
    print("hotspot:", len(h))
