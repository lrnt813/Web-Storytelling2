"""Ekspor seluruh tabel Bab IV dari hasil TERKUNCI (data/locked/) ke output_bab4/.

    python -m scripts.export_bab4                  # K utama dari pengaturan_hasil.json (bawaan 3)
    python -m scripts.export_bab4 --k-utama 2      # K utama lain (K lainnya menjadi sensitivitas)
    python -m scripts.export_bab4 --keluaran DIR   # tulis ke folder lain (mis. uji)

Keluaran:
  T*.csv   tabel utama Bab IV (K utama untuk tabel yang bergantung K)
  L*.csv   tabel lampiran (metrik pendukung pemilihan K)
  S*.csv   tabel sensitivitas K lainnya, struktur identik dengan tabel T bernomor sama
  Tabel_Bab4.docx          semua tabel, format angka Indonesia (koma desimal, titik ribuan;
                           waktu 2 desimal, metrik 3 desimal)
  ringkasan_angka_bab4.md  angka kunci kedua K (K utama di atas)
  perbandingan_v2_vs_v3.md angka kunci v2 (data/locked/arsip_v2/) vs v3 dan penyebab perubahan
  peta/                    peta tipologi per level per K dan peta TAS per level (PNG)

Skrip ini hanya membaca data/locked/ (termasuk arsip_v2/ untuk perbandingan), pengaturan_hasil.json,
dan geometri statis grid (frontend/static_grid_kulonprogo.geojson) untuk peta.
Folder arsip di output_bab4/ tidak disentuh.
"""
import argparse
import hashlib
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parents[1]
LOCKED = PROJECT_ROOT / "data" / "locked"
ARSIP_V2 = LOCKED / "arsip_v2" / "thesis_results.json"
OUT = PROJECT_ROOT / "output_bab4"
PENGATURAN = PROJECT_ROOT / "pengaturan_hasil.json"
GEOM = PROJECT_ROOT / "frontend" / "static_grid_kulonprogo.geojson"

KAT = ["pendidikan", "kesehatan", "pemerintahan", "ibadah", "gor"]
KAT_LABEL = {"pendidikan": "Pendidikan", "kesehatan": "Kesehatan", "pemerintahan": "Pemerintahan",
             "ibadah": "Tempat Ibadah", "gor": "GOR/Gedung Serbaguna"}
LEVEL_ORDER = ["baseline", "rendah", "sedang", "tinggi"]
LEVEL_LABEL = {"baseline": "Baseline", "rendah": "Rendah", "sedang": "Sedang", "tinggi": "Tinggi"}
TIME_COLS = [f"waktu_tes_{k}" for k in KAT] + ["waktu_tes_min"]
VAR_LABEL = {**{f"waktu_tes_{k}": f"Waktu TES {v} (menit)" for k, v in KAT_LABEL.items()},
             "waktu_tes_min": "Waktu TES minimum (menit)"}
CLUSTER_COLORS = ["#1f77b4", "#2ca02c", "#9467bd", "#e377c2", "#bcbd22", "#17becf", "#8c564b", "#d62728",
                  "#7f7f7f", "#ff7f0e"]
TERGENANG_COLOR = "#38bdf8"


# ── format angka Indonesia ──────────────────────────────────────────────────
def id_num(v, dec=2):
    if v is None or (isinstance(v, float) and v != v):
        return "–"
    if isinstance(v, bool):
        return "Ya" if v else "Tidak"
    if isinstance(v, (int, np.integer)) and dec == 0:
        return f"{int(v):,}".replace(",", ".")
    s = f"{float(v):,.{dec}f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


class Table:
    """Tabel dengan data mentah (untuk CSV) dan format kolom (untuk DOCX)."""

    def __init__(self, code, title, df, fmt=None, note=""):
        self.code, self.title, self.df, self.fmt, self.note = code, title, df, fmt or {}, note

    def formatted(self):
        out = self.df.copy().astype(object)
        for col in out.columns:
            dec = self.fmt.get(col)
            if dec is not None:
                out[col] = [id_num(v, dec) for v in self.df[col]]
            else:
                out[col] = ["–" if (v is None or (isinstance(v, float) and v != v)) else str(v)
                            for v in self.df[col]]
        return out


def load_locked():
    lock = json.loads((LOCKED / "LOCK.json").read_text(encoding="utf-8"))
    res_path = LOCKED / "thesis_results.json"
    sha = hashlib.sha256(res_path.read_bytes()).hexdigest()
    if lock.get("files", {}).get("thesis_results.json") != sha:
        raise SystemExit("thesis_results.json di data/locked/ tidak cocok dengan checksum LOCK.json.")
    grid = pd.read_csv(LOCKED / "thesis_grid_results.csv.gz")
    return json.loads(res_path.read_text(encoding="utf-8")), lock, grid


def k_utama_default() -> int:
    try:
        return int(json.loads(PENGATURAN.read_text(encoding="utf-8"))["k_utama"])
    except (OSError, KeyError, ValueError):
        return 3


def state_name(s, k):
    return "Tergenang" if s == k else f"K{s}"


# ── tabel tidak bergantung K ────────────────────────────────────────────────
def tables_common(R):
    T = []
    lv, ds, pool = R["levels"], R["data_summary"], R["data_gabungan"]
    rows = [("Grid analisis 100 × 100 m", ds["n_grid"]), ("Segmen jaringan jalan", ds["n_road_segments"])]
    rows += [(f"TES {KAT_LABEL[c]}", ds["tes_counts"][c]) for c in KAT]
    rows += [("TES total", sum(ds["tes_counts"].values())),
             ("Baris data gabungan (grid × level non-Tergenang)", pool["n_baris"])]
    T.append(Table("T01", "Ringkasan data spasial", pd.DataFrame(rows, columns=["Komponen", "Jumlah"]),
                   {"Jumlah": 0}, f"Sistem koordinat {ds['crs']}; kecepatan berjalan "
                   f"{id_num(R['meta']['walking_speed_m_per_min'], 0)} m/menit; waktu penalti T_pen = 3 × T_max "
                   f"Baseline = {id_num(R['t_pen'], 2)} menit. Kelas bahaya banjir grid, ruas, dan TES diturunkan dari "
                   f"raster InaRisk (grid: mayoritas piksel; ruas: maksimum; TES: piksel di titik)."))

    rows = []
    for key in LEVEL_ORDER:
        d = lv[key]["diagnostik"]
        rows.append({"Level": LEVEL_LABEL[key], "Kelas bahaya ditutup": ", ".join(map(str, lv[key]["kelas_ditutup"])) or "tidak ada",
                     **{f"TES {KAT_LABEL[c]}": d["tes_valid_per_kategori"][c] for c in KAT},
                     "TES valid total": d["tes_valid_total"], "Ruas jalan ditutup": d["ruas_jalan_ditutup"],
                     "Baris data gabungan": pool["n_baris_per_level"][key]})
    df = pd.DataFrame(rows)
    T.append(Table("T02", "Definisi level, TES valid, dan ruas jalan ditutup", df,
                   {c: 0 for c in df.columns if c not in ("Level", "Kelas bahaya ditutup")}))

    rows = []
    for key in LEVEL_ORDER:
        s, d = lv[key], lv[key]["diagnostik"]
        rows.append({"Level": LEVEL_LABEL[key], "Grid Tergenang": s["n_grid_tergenang"],
                     "% grid": s["persen_grid_tergenang"], "Grid non-Tergenang": s["n_grid_non_tergenang"],
                     "Tergenang & terisolasi": d["grid_terisolasi_tergenang"],
                     "Rerata waktu minimum grid Tergenang (menit, informasi)": s["mean_waktu_min_tergenang"]})
    df = pd.DataFrame(rows)
    T.append(Table("T03", "Grid Tergenang per level", df,
                   {"Grid Tergenang": 0, "% grid": 2, "Grid non-Tergenang": 0, "Tergenang & terisolasi": 0,
                    "Rerata waktu minimum grid Tergenang (menit, informasi)": 2},
                   "Grid Tergenang dikeluarkan dari klasterisasi dan TAS; waktu tempuhnya hanya informasi."))

    rows = [{"Variabel": VAR_LABEL[s["variabel"]], "n": s["jumlah"], "Rata-rata": s["mean"],
             "Simpangan baku": s["std"], "Min": s["min"], "Q1": s["q25"], "Median": s["median"],
             "Q3": s["q75"], "Maks": s["max"]} for s in R["baseline_time_stats"]]
    df = pd.DataFrame(rows)
    T.append(Table("T04", "Statistik deskriptif waktu tempuh ke TES pada Baseline (menit)", df,
                   {"n": 0} | {c: 2 for c in df.columns if c not in ("Variabel", "n")},
                   "Waktu = (ruas snapping centroid → jalan + jarak jaringan + ruas snapping jalan → TES) / 80 m/menit."))

    ks = R["k_selection"]
    best = ks["ari_subsampel_tertinggi"]
    rows = [{"K": c["k"], "ARI subsampel (rerata)": c["ari_subsampel_mean"], "ARI subsampel (sd)": c["ari_subsampel_sd"],
             "ARI inisialisasi (rerata)": c["ari_inisialisasi_mean"], "ARI inisialisasi (sd)": c["ari_inisialisasi_sd"],
             "Silhouette": c["silhouette"], "Ketegasan partisi": c["ketegasan_partisi"],
             "Size entropy": c["size_entropy"], "Klaster terbesar (%)": c["klaster_terbesar_persen"],
             "Setara secara stabilitas dengan K terbaik": bool(best - c["ari_subsampel_mean"] <= ks["toleransi"] + 1e-12)}
            for c in ks["candidates"]]
    df = pd.DataFrame(rows)
    T.append(Table("T05", "Pemilihan K berbasis stabilitas (data gabungan)", df,
                   {c: 3 for c in df.columns if c not in ("K", "Klaster terbesar (%)",
                                                         "Setara secara stabilitas dengan K terbaik")}
                   | {"K": 0, "Klaster terbesar (%)": 2, "Setara secara stabilitas dengan K terbaik": 0},
                   f"{ks['B']} subsampel {id_num(ks['fraksi_subsampel'] * 100, 0)} % id_grid, {ks['n_init_subsampel']} "
                   f"inisialisasi per subsampel; ARI inisialisasi = rerata 45 pasangan run seed 42–51. \"Setara\" = selisih "
                   f"rerata ARI subsampel dari nilai tertinggi ({id_num(best, 3)}) ≤ {id_num(ks['toleransi'], 2)}. Aturan "
                   f"stabilitas dengan pemecah seri \"K terkecil\" memilih K = {ks['k_terpilih']}. Model utama ditetapkan "
                   f"peneliti bersama pembimbing (lihat METODOLOGI §9)."))
    rows = [{"K": c["k"], "I-Index": c["iidx"], "Dunn (rerata)": c["dunn"], "Dunn (sd)": c["dunn_sd"],
             "DESC": c["desc"], "PESC": c["pesc"], "Komposit +DESC": c["komposit_dengan_desc"],
             "Komposit −DESC": c["komposit_tanpa_desc"]} for c in ks["candidates"]]
    df = pd.DataFrame(rows)
    T.append(Table("L01", "Lampiran: metrik pendukung pemilihan K (tidak dipakai memilih)", df,
                   {"K": 0, "I-Index": 3, "Dunn (rerata)": 4, "Dunn (sd)": 4, "DESC": 3, "PESC": 4,
                    "Komposit +DESC": 3, "Komposit −DESC": 3},
                   "DESC/PESC: blok rook; Dunn: 5 sampel 2.000 baris (seed 42–46). PESC tidak stabil antar-K."))

    rows, rows_s, rows_d, rows_k = [], [], [], []
    for key in LEVEL_ORDER:
        s = lv[key]["tas"]
        sp = s["sensitivitas_persentil"]
        rows.append({"Level": LEVEL_LABEL[key], "TAS": s["jumlah_tas"], "Non-TAS": s["jumlah_non_tas"],
                     "Terputus": s["jumlah_terputus"], "Tergenang": s["jumlah_tergenang"],
                     "% TAS (semua grid)": s["persen_tas"], "% TAS (non-Tergenang)": s["persen_tas_non_tergenang"],
                     "Median T_ideal TAS (menit)": s["median_t_ideal_tas"],
                     "Median T_aktual TAS (menit)": s["median_t_aktual_tas"]})
        rows_s.append({"Level": LEVEL_LABEL[key], "TAS aturan utama": s["jumlah_tas"],
                       "TAS aturan persentil": sp["jumlah_tas"], "P25 T_ideal (menit)": sp["p25_t_ideal"],
                       "P75 DI": sp["p75_di"], "TAS kedua aturan": sp["jumlah_tas_kedua_aturan"],
                       "Hanya aturan utama": sp["jumlah_tas_hanya_aturan_utama"],
                       "Hanya aturan persentil": sp["jumlah_tas_hanya_persentil"]})
        for grp, q in (("TAS", s["sebaran_di_tas"]), ("Non-TAS", s["sebaran_di_non_tas"])):
            rows_d.append({"Level": LEVEL_LABEL[key], "Kelompok": grp, "n": q["n"], "Median DI": q["median"],
                           "P25": q["p25"], "P75": q["p75"], "P90": q["p90"]})
        rows_k.append({"Level": LEVEL_LABEL[key], **{KAT_LABEL[c]: s["tas_per_kategori_tes_terdekat"].get(c, 0) for c in KAT}})
    df = pd.DataFrame(rows)
    T.append(Table("T11", "Titik Aman Semu per level (aturan utama: DI ≥ 2 dan T_ideal ≤ 5 menit)", df,
                   {c: 0 for c in ("TAS", "Non-TAS", "Terputus", "Tergenang")}
                   | {"% TAS (semua grid)": 2, "% TAS (non-Tergenang)": 2, "Median T_ideal TAS (menit)": 2,
                      "Median T_aktual TAS (menit)": 2},
                   "Grid non-Tergenang yang terjangkau: TAS bila DI_t ≥ 2 dan T_ideal ≤ 5 menit (garis lurus ≤ 400 m). "
                   "Batas minimum 50 m pada T_ideal dan T_aktual. Terputus = non-Tergenang tidak menjangkau TES."))
    df = pd.DataFrame(rows_s)
    T.append(Table("T11b", "Sensitivitas TAS: aturan persentil (T_ideal ≤ P25 dan DI ≥ P75)", df,
                   {c: 0 for c in df.columns if c not in ("Level", "P25 T_ideal (menit)", "P75 DI")}
                   | {"P25 T_ideal (menit)": 2, "P75 DI": 3}))
    df = pd.DataFrame(rows_d)
    T.append(Table("T11c", "Sebaran Detour Index pada TAS dan Non-TAS (aturan utama)", df,
                   {"n": 0, "Median DI": 3, "P25": 3, "P75": 3, "P90": 3},
                   "Deskriptif. Perbedaan sebaran DI antara TAS dan Non-TAS mengikuti definisi aturan, sehingga tidak "
                   "dapat dipakai sebagai bukti keberhasilan deteksi."))
    df = pd.DataFrame(rows_k)
    T.append(Table("T11d", "Jumlah TAS per kategori TES terdekat (aturan utama)", df,
                   {KAT_LABEL[c]: 0 for c in KAT}, "TES valid terdekat secara Euclidean."))

    rows = []
    for key in LEVEL_ORDER:
        d = lv[key]["diagnostik"]
        kb = d["kelas_bahaya"]
        rows.append({"Level": LEVEL_LABEL[key], "Grid Tergenang": d["grid_tergenang"],
                     "Terisolasi (non-Tergenang)": d["grid_terisolasi_non_tergenang"],
                     "Terisolasi (Tergenang)": d["grid_terisolasi_tergenang"],
                     "Ruas ditutup": d["ruas_jalan_ditutup"],
                     "Kelas ditutup (aturan)": ", ".join(map(str, kb["kelas_ditutup_aturan"])) or "–",
                     "Kelas grid Tergenang (data)": ", ".join(map(str, kb["grid_tergenang_kelas_data"])) or "–",
                     "Kelas ruas ditutup (data)": ", ".join(map(str, kb["ruas_ditutup_kelas_data"])) or "–",
                     "Aturan = data": kb["konsisten"],
                     "TAS: T_aktual < T_ideal": lv[key]["tas"].get("jumlah_t_aktual_lt_t_ideal")})
    df = pd.DataFrame(rows)
    T.append(Table("T12", "Diagnostik per level", df,
                   {"Grid Tergenang": 0, "Terisolasi (non-Tergenang)": 0, "Terisolasi (Tergenang)": 0,
                    "Ruas ditutup": 0, "Aturan = data": 0, "TAS: T_aktual < T_ideal": 0},
                   "Terisolasi = jumlah_opsi_rute = 0. Kolom terakhir harus 0."))

    rows = [{"Level": LEVEL_LABEL[key], "Rata-rata waktu minimum": lv[key]["mean_waktu_min"],
             "Median waktu minimum": lv[key]["median_waktu_min"],
             **{f"Rata-rata {KAT_LABEL[c]}": lv[key]["mean_waktu"][c] for c in KAT}} for key in LEVEL_ORDER]
    df = pd.DataFrame(rows)
    T.append(Table("T13", "Rata-rata waktu tempuh ke TES per level, grid non-Tergenang (menit)", df,
                   {c: 2 for c in df.columns if c != "Level"}))

    pp = R["preprocessing"]
    rows = [{"Tahap": "Data fit", "Nilai": f"data gabungan, {id_num(pp['n_baris_fit'], 0)} baris"},
            {"Tahap": "Fitur masukan", "Nilai": ", ".join(pp["features_in"])},
            {"Tahap": "Fitur terpilih", "Nilai": ", ".join(pp["features_kept"])},
            {"Tahap": "Fitur dibuang", "Nilai": ", ".join(pp["features_dropped"]) or "–"},
            {"Tahap": "Capping 3 × IQR", "Nilai": "; ".join(f"{c}: [{id_num(a, 3)}; {id_num(b, 3)}]"
                                                           for c, (a, b) in pp["caps"].items())},
            {"Tahap": "Komponen PCA", "Nilai": id_num(pp["n_components"], 0)},
            {"Tahap": "Varians kumulatif PCA", "Nilai": id_num(pp["explained_variance"] * 100, 2) + " %"},
            {"Tahap": "σ kernel Gaussian per level (m)",
             "Nilai": "; ".join(f"{LEVEL_LABEL[k2]} {id_num(v, 2)}" for k2, v in pool["sigma_per_level"].items())}]
    T.append(Table("T14", "Pra-pemrosesan (di-fit sekali pada data gabungan)", pd.DataFrame(rows)))
    return T


# ── tabel bergantung K ──────────────────────────────────────────────────────
def tables_for_k(R, k, prefix):
    T = []
    M = R["model"][str(k)]
    sfx = f"K{k}"

    if M.get("algorithm_comparison"):
        rows = []
        for a in M["algorithm_comparison"]:
            rows.append({"Algoritma": a["algoritma"],
                         "Status": {"ok": "layak", "degeneratif": "degeneratif — tidak layak dibandingkan",
                                    "gagal": "gagal — tidak layak dibandingkan"}[a["status"]],
                         "m": a.get("m"), "Inisialisasi": a["n_init"], "Jumlah klaster": a.get("n_cluster"),
                         "Silhouette": a.get("silhouette"), "Calinski-Harabasz": a.get("calinski_harabasz"),
                         "Davies-Bouldin": a.get("davies_bouldin"), "Moran's I": a.get("moran_i"),
                         "p (Moran)": a.get("moran_p"), "Proporsi tetangga sama": a.get("proporsi_tetangga_sama"),
                         "PC": a.get("pc"), "PE": a.get("pe"),
                         "Size entropy": a.get("size_entropy"), "Klaster terbesar (%)": a.get("klaster_terbesar_persen"),
                         "Waktu per inisialisasi (detik)": a["time_per_init_sec"], "Galat": a.get("galat", "")})
        df = pd.DataFrame(rows)
        T.append(Table(f"{prefix}06", f"Perbandingan algoritma pada Baseline {sfx}", df,
                       {"m": 1, "Inisialisasi": 0, "Jumlah klaster": 0, "Silhouette": 3, "Calinski-Harabasz": 3,
                        "Davies-Bouldin": 3, "Moran's I": 3, "p (Moran)": 3, "Proporsi tetangga sama": 3, "PC": 3,
                        "PE": 3, "Size entropy": 3, "Klaster terbesar (%)": 2, "Waktu per inisialisasi (detik)": 2},
                       "FCM (α = 0), SFCM, SDWFCM: m = 1,7, 10 inisialisasi (J terkecil), praproses dan penomoran sama. "
                       "PC/PE dari keanggotaan akhir. Moran's I label: 999 permutasi. Degeneratif = klaster terbesar > 90 %."))

    rows, rows_t = [], []
    for p in M["cluster_profile"]:
        rows.append({"Klaster": p["klaster"], "Jumlah baris": p["jumlah_baris"], "% baris": p["persen_baris"],
                     "Median waktu minimum": p["waktu_tes_min_median"], "P25 waktu minimum": p["waktu_tes_min_p25"],
                     "P75 waktu minimum": p["waktu_tes_min_p75"], "Rerata waktu minimum": p["waktu_tes_min"],
                     "Baris waktu minimum = penalti": p["waktu_tes_min_n_penalti"],
                     "% baris waktu minimum = penalti": p["waktu_tes_min_persen_penalti"],
                     "Proporsi terisolasi": p["proporsi_terisolasi"], "Opsi rute (rerata)": p["jumlah_opsi_rute"],
                     "Kepadatan jalan (rerata)": p["Road_Density_mean"], "Kelas bahaya (deskriptif)": p["banjir"],
                     "Interpretasi": p["deskripsi"]})
        for c in TIME_COLS:
            rows_t.append({"Klaster": p["klaster"], "Variabel": VAR_LABEL[c], "Rerata": p[c],
                           "Median": p[f"{c}_median"], "P25": p[f"{c}_p25"], "P75": p[f"{c}_p75"],
                           "Baris penalti": p[f"{c}_n_penalti"], "% baris penalti": p[f"{c}_persen_penalti"]})
    df = pd.DataFrame(rows)
    T.append(Table(f"{prefix}07", f"Profil tipologi SDWFCM {sfx} (data gabungan keempat level)", df,
                   {"Klaster": 0, "Jumlah baris": 0, "% baris": 2, "Median waktu minimum": 2, "P25 waktu minimum": 2,
                    "P75 waktu minimum": 2, "Rerata waktu minimum": 2, "Baris waktu minimum = penalti": 0,
                    "% baris waktu minimum = penalti": 2, "Proporsi terisolasi": 3, "Opsi rute (rerata)": 2,
                    "Kepadatan jalan (rerata)": 3, "Kelas bahaya (deskriptif)": 3},
                   "Label: \"Terisolasi\" bila > 50 % baris waktu minimum = penalti; selain itu median waktu minimum "
                   "≤ 10 menit \"Akses Baik\", 10–30 \"Akses Sedang\", > 30 \"Akses Kritis\". Klaster 0 = rerata waktu "
                   "minimum terkecil. Kelas bahaya bukan fitur."))
    df = pd.DataFrame(rows_t)
    T.append(Table(f"{prefix}07b", f"Waktu tempuh per kategori TES per klaster {sfx} (menit)", df,
                   {"Klaster": 0, "Rerata": 2, "Median": 2, "P25": 2, "P75": 2, "Baris penalti": 0, "% baris penalti": 2}))

    rows = []
    for key in LEVEL_ORDER:
        row = {"Level": LEVEL_LABEL[key]}
        for d in M["distribusi"][key]:
            nm = "Tergenang" if d["nama"] == "Tergenang" else f"K{d['state']}"
            row[f"{nm} (grid)"] = d["jumlah_grid"]
            row[f"{nm} (%)"] = d["persen_semua_grid"]
        rows.append(row)
    df = pd.DataFrame(rows)
    T.append(Table(f"{prefix}08", f"Distribusi tipologi dan Tergenang per level {sfx}", df,
                   {c: (0 if c.endswith("(grid)") else 2) for c in df.columns if c != "Level"},
                   "Persentase terhadap seluruh grid."))

    for i, t in enumerate(M["transitions"]):
        S = k + 1
        mat = pd.DataFrame(t["matrix"], columns=[f"ke {state_name(j, k)}" for j in range(S)])
        mat.insert(0, "Dari", [state_name(j, k) for j in range(S)])
        T.append(Table(f"{prefix}09{chr(97 + i)}", f"Matriks transisi {LEVEL_LABEL[t['from_level']]} → "
                       f"{LEVEL_LABEL[t['to_level']]} {sfx} (jumlah grid)", mat, {c: 0 for c in mat.columns if c != "Dari"}))
    rows = []
    for t in M["transitions"]:
        d = t["dominant_transition"]
        rows.append({"Transisi": f"{LEVEL_LABEL[t['from_level']]} → {LEVEL_LABEL[t['to_level']]}",
                     "Masuk Tergenang (grid)": t["masuk_tergenang"],
                     "Masuk Tergenang (% semua grid)": t["persen_masuk_tergenang_semua_grid"],
                     "Grid non-Tergenang di kedua level": t["n_non_tergenang_kedua_level"],
                     "Stability rate (%)": t["stability_rate"], "ARI": t["ari"],
                     "Transisi dominan": f"{state_name(d['from'], k)} → {state_name(d['to'], k)}" if d else "–",
                     "Jumlah (dominan)": d["count"] if d else None, "Dominan menuju": d["menuju"] if d else "–",
                     "Active edges": t["active_edges"], "Edge mungkin": (k + 1) * k, "CDVM": t["cdvm"]})
    df = pd.DataFrame(rows)
    T.append(Table(f"{prefix}10", f"Ringkasan transisi antarlevel {sfx} (state tipologi + Tergenang)", df,
                   {"Masuk Tergenang (grid)": 0, "Masuk Tergenang (% semua grid)": 2,
                    "Grid non-Tergenang di kedua level": 0, "Stability rate (%)": 2, "ARI": 3, "Jumlah (dominan)": 0,
                    "Active edges": 0, "Edge mungkin": 0, "CDVM": 3},
                   "Stability rate dan ARI hanya pada grid non-Tergenang di kedua level. CDVM = ½ Σ |p_s(tujuan) − p_s(asal)|."))

    v = M["model_final"]["validity"]
    rows = [{"Metrik": "I-Index", "Nilai": v["iidx"]}, {"Metrik": "Ketegasan partisi", "Nilai": v["ketegasan_partisi"]},
            {"Metrik": "Koefisien partisi (PC)", "Nilai": v["pc"]}, {"Metrik": "Entropi partisi (PE)", "Nilai": v["pe"]},
            {"Metrik": "Dunn (rerata 5 sampel)", "Nilai": v["dunn"]["mean"]},
            {"Metrik": "Dunn (sd 5 sampel)", "Nilai": v["dunn"]["sd"]},
            {"Metrik": "DESC (blok rook)", "Nilai": v["desc"]}, {"Metrik": "PESC (km)", "Nilai": v["pesc"]},
            {"Metrik": "Silhouette (sampel 10.000)", "Nilai": v["silhouette_sampel"]},
            {"Metrik": "Size entropy", "Nilai": v["size_entropy"]},
            {"Metrik": "Klaster terbesar (%)", "Nilai": v["klaster_terbesar_persen"]},
            {"Metrik": "Proporsi tetangga rook berlabel sama", "Nilai": v["proporsi_tetangga_sama"]}]
    T.append(Table(f"{prefix}15", f"Validitas model final {sfx} (data gabungan)", pd.DataFrame(rows), {"Nilai": 4},
                   f"Solusi J terkecil dari 10 inisialisasi: seed {M['model_final']['seed_terbaik']}, "
                   f"J = {id_num(M['model_final']['objective'], 3)}."))
    return T


# ── perbandingan K = 2 vs K = 3 ─────────────────────────────────────────────
def worst_drop(grid, k, a="sedang", b="tinggi"):
    """Grid yang non-Tergenang di kedua level, bukan tipologi terburuk (K−1) di a, dan tipologi terburuk di b."""
    sa, sb = grid[f"state_{a}_k{k}"].values, grid[f"state_{b}_k{k}"].values
    return int(((sa < k - 1) & (sb == k - 1)).sum())


def tables_k2_vs_k3(R, grid):
    T = []
    ks = [int(k) for k in R["k_keluaran"]]
    rows = []
    for i, pair in enumerate(R["model"][str(ks[0])]["transitions"]):
        row = {"Transisi": f"{LEVEL_LABEL[pair['from_level']]} → {LEVEL_LABEL[pair['to_level']]}"}
        for k in ks:
            t = R["model"][str(k)]["transitions"][i]
            d = t["dominant_transition"]
            row[f"Dominan K{k}"] = (f"{state_name(d['from'], k)} → {state_name(d['to'], k)} ({id_num(d['count'], 0)})"
                                    if d else "–")
            row[f"SR K{k} (%)"] = t["stability_rate"]
            row[f"ARI K{k}"] = t["ari"]
        rows.append(row)
    df = pd.DataFrame(rows)
    T.append(Table("T16", "Temuan utama K = 2 vs K = 3: transisi", df,
                   {c: (2 if c.startswith("SR") else 3) for c in df.columns if c.startswith(("SR", "ARI"))}))

    rows = []
    for key in LEVEL_ORDER[1:]:
        row = {"Level vs Baseline": LEVEL_LABEL[key]}
        for k in ks:
            base = {d["state"]: d["persen_semua_grid"] for d in R["model"][str(k)]["distribusi"]["baseline"]}
            cur = {d["state"]: d["persen_semua_grid"] for d in R["model"][str(k)]["distribusi"][key]}
            row[f"Perubahan K{k} (poin persen)"] = "; ".join(
                f"{state_name(s, k)} {'+' if cur[s] - base[s] >= 0 else '−'}{id_num(abs(cur[s] - base[s]), 1)}"
                for s in sorted(cur))
        rows.append(row)
    df = pd.DataFrame(rows)
    T.append(Table("T16b", "Temuan utama K = 2 vs K = 3: arah perubahan distribusi per level", df))

    rows = [{"Ukuran": "Grid yang turun ke tipologi terburuk (K−1) pada Sedang → Tinggi",
             **{f"K{k}": worst_drop(grid, k) for k in ks}},
            {"Ukuran": "Tipologi terburuk (label interpretasi)",
             **{f"K{k}": R["model"][str(k)]["cluster_names"][str(k - 1)] for k in ks}}]
    T.append(Table("T16c", "Temuan utama K = 2 vs K = 3: penurunan ke tipologi terburuk", pd.DataFrame(rows),
                   {f"K{k}": None for k in ks},
                   "Dihitung pada grid non-Tergenang di Sedang dan Tinggi yang di Sedang bukan tipologi terburuk."))

    if R.get("tabulasi_silang_k2_k3"):
        ct = R["tabulasi_silang_k2_k3"]
        df = pd.DataFrame(ct["matrix"], columns=ct["kolom"])
        df.insert(0, "K=2 \\ K=3", ct["baris"])
        tot = df[ct["kolom"]].sum(axis=1)
        for c in ct["kolom"]:
            df[f"{c} (% baris K2)"] = df[c] / tot * 100
        T.append(Table("T17", "Tabulasi silang label K = 2 × K = 3 (data gabungan)", df,
                       {c: (0 if "(%" not in c else 2) for c in df.columns if c != "K=2 \\ K=3"},
                       f"Baris = klaster K = 2, kolom = klaster K = 3; ARI antar-partisi = {id_num(ct['ari'], 3)}."))
    return T


# ── peta ────────────────────────────────────────────────────────────────────
def write_maps(R, grid, out_dir):
    import geopandas as gpd
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Patch
    if not GEOM.exists():
        return
    g = gpd.read_file(GEOM)[["id_grid", "geometry"]].merge(grid, on="id_grid", how="left")
    d = out_dir / "peta"
    d.mkdir(exist_ok=True)
    for k in [int(x) for x in R["k_keluaran"]]:
        names = R["model"][str(k)]["cluster_names"]
        for key in LEVEL_ORDER:
            st = g[f"state_{key}_k{k}"].values
            colors = np.where(st == k, TERGENANG_COLOR, np.array(CLUSTER_COLORS)[np.minimum(st, 9)])
            fig, ax = plt.subplots(figsize=(7, 8.5), dpi=120)
            g.plot(ax=ax, color=colors, linewidth=0)
            ax.set_axis_off()
            ax.set_title(f"Tipologi K = {k} — {R['levels'][key]['label_lengkap']}", fontsize=9)
            ax.legend(handles=[Patch(color=CLUSTER_COLORS[c], label=f"K{c}: {names[str(c)]}") for c in range(k)]
                      + [Patch(color=TERGENANG_COLOR, label="Tergenang")], loc="lower left", fontsize=7)
            fig.tight_layout()
            fig.savefig(d / f"tipologi_K{k}_{key}.png")
            plt.close(fig)
    pal = {0: "#e5e7eb", 1: "#a855f7", 2: "#0f172a", 3: TERGENANG_COLOR}
    lab = {0: "Non-TAS", 1: "TAS", 2: "Terputus", 3: "Tergenang"}
    for key in LEVEL_ORDER:
        st = g[f"tas_status_{key}"].values
        fig, ax = plt.subplots(figsize=(7, 8.5), dpi=120)
        g.plot(ax=ax, color=[pal[int(s)] for s in st], linewidth=0)
        ax.set_axis_off()
        ax.set_title(f"Titik Aman Semu (DI ≥ 2, T_ideal ≤ 5 menit) — {R['levels'][key]['label_lengkap']}", fontsize=9)
        ax.legend(handles=[Patch(color=pal[s], label=f"{lab[s]} ({id_num(int((st == s).sum()), 0)})") for s in pal],
                  loc="lower left", fontsize=7)
        fig.tight_layout()
        fig.savefig(d / f"tas_{key}.png")
        plt.close(fig)


# ── keluaran ────────────────────────────────────────────────────────────────
def _slug(s):
    import re
    return re.sub(r"[^0-9A-Za-z]+", "_", s).strip("_")[:70]


def write_csv(tables, out_dir):
    for old in list(out_dir.glob("T*.csv")) + list(out_dir.glob("S*.csv")) + list(out_dir.glob("L*.csv")):
        old.unlink()
    for t in tables:
        t.df.to_csv(out_dir / f"{t.code}_{_slug(t.title)}.csv", index=False, encoding="utf-8-sig")


def write_docx(tables, lock, k_utama, out_dir):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.5))
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9)
    doc.add_heading("Tabel Bab IV — hasil terkunci", level=1)
    doc.add_paragraph(f"Sumber: data/locked/ (commit analisis {lock['git']['commit_analisis'][:10]}, "
                      f"dikunci {lock['locked_at']}). K utama = {k_utama}; tabel berawalan S = sensitivitas K lain; "
                      f"L = lampiran. Format angka Indonesia; waktu 2 desimal, metrik 3 desimal.")
    for t in tables:
        doc.add_heading(f"{t.code}. {t.title}", level=2)
        f = t.formatted()
        table = doc.add_table(rows=1, cols=len(f.columns))
        table.style = "Table Grid"
        for j, col in enumerate(f.columns):
            cell = table.rows[0].cells[j]
            cell.text = str(col)
            for r in cell.paragraphs[0].runs:
                r.font.bold = True
                r.font.size = Pt(8)
        for _, row in f.iterrows():
            cells = table.add_row().cells
            for j, val in enumerate(row):
                cells[j].text = str(val)
                for r in cells[j].paragraphs[0].runs:
                    r.font.size = Pt(8)
        if t.note:
            p = doc.add_paragraph(t.note)
            p.runs[0].font.size = Pt(8)
            p.runs[0].font.italic = True
    doc.save(out_dir / "Tabel_Bab4.docx")


def _summary_k(R, k, grid):
    M = R["model"][str(k)]
    L = [f"### Model K = {k}", ""]
    v = M["model_final"]["validity"]
    L.append(f"- Validitas: Silhouette (sampel) {id_num(v['silhouette_sampel'], 3)}, ketegasan partisi "
             f"{id_num(v['ketegasan_partisi'], 3)}, PC {id_num(v['pc'], 3)}, PE {id_num(v['pe'], 3)}, klaster terbesar "
             f"{id_num(v['klaster_terbesar_persen'], 2)} %")
    for p in M["cluster_profile"]:
        L.append(f"- K{p['klaster']} — **{p['deskripsi']}**: {id_num(p['jumlah_baris'], 0)} baris ({id_num(p['persen_baris'], 2)} %); "
                 f"waktu minimum median {id_num(p['waktu_tes_min_median'], 2)} [P25 {id_num(p['waktu_tes_min_p25'], 2)}; "
                 f"P75 {id_num(p['waktu_tes_min_p75'], 2)}] menit, rerata {id_num(p['waktu_tes_min'], 2)}; penalti "
                 f"{id_num(p['waktu_tes_min_persen_penalti'], 2)} %; terisolasi {id_num(p['proporsi_terisolasi'], 3)}")
    for key in LEVEL_ORDER:
        dist = ", ".join(f"{state_name(x['state'], k)} = {id_num(x['jumlah_grid'], 0)} ({id_num(x['persen_semua_grid'], 2)} %)"
                         for x in M["distribusi"][key])
        L.append(f"- {LEVEL_LABEL[key]}: {dist}")
    for t in M["transitions"]:
        d = t["dominant_transition"]
        L.append(f"- {LEVEL_LABEL[t['from_level']]} → {LEVEL_LABEL[t['to_level']]}: masuk Tergenang "
                 f"{id_num(t['masuk_tergenang'], 0)}; SR {id_num(t['stability_rate'], 2)} %; ARI {id_num(t['ari'], 3)}; "
                 f"dominan {state_name(d['from'], k)} → {state_name(d['to'], k)} ({id_num(d['count'], 0)}); CDVM {id_num(t['cdvm'], 3)}")
    L.append(f"- Grid turun ke tipologi terburuk Sedang → Tinggi: {id_num(worst_drop(grid, k), 0)}")
    for a in M.get("algorithm_comparison", []):
        if a["status"] == "gagal":
            L.append(f"- {a['algoritma']}: **gagal** — {a['galat']}")
        else:
            L.append(f"- {a['algoritma']}{'' if a['status'] == 'ok' else ' (**degeneratif**)'}: Silhouette "
                     f"{id_num(a.get('silhouette'), 3)}, CH {id_num(a.get('calinski_harabasz'), 1)}, DB {id_num(a.get('davies_bouldin'), 3)}, "
                     f"Moran's I {id_num(a.get('moran_i'), 3)}, PC {id_num(a.get('pc'), 3)}, PE {id_num(a.get('pe'), 3)}, "
                     f"klaster terbesar {id_num(a.get('klaster_terbesar_persen'), 2)} %")
    return L + [""]


def write_summary(R, lock, k_utama, grid, out_dir):
    lv, ks = R["levels"], R["k_selection"]
    order = [k_utama] + [int(k) for k in R["k_keluaran"] if int(k) != k_utama]
    L = ["# Ringkasan angka kunci Bab IV", "",
         f"Sumber: `data/locked/` — commit analisis `{lock['git']['commit_analisis']}`, dikunci {lock['locked_at']}. "
         f"K utama = **{k_utama}** (pengaturan_hasil.json); K lain = sensitivitas.", "",
         "## Data dan level", "",
         f"- Grid {id_num(R['data_summary']['n_grid'], 0)}; segmen jalan {id_num(R['data_summary']['n_road_segments'], 0)}; "
         f"TES {id_num(sum(R['data_summary']['tes_counts'].values()), 0)}; T_pen {id_num(R['t_pen'], 2)} menit",
         f"- Data gabungan {id_num(R['data_gabungan']['n_baris'], 0)} baris"]
    for key in LEVEL_ORDER:
        s, t = lv[key], lv[key]["tas"]
        L.append(f"- {s['label_lengkap']}: Tergenang {id_num(s['n_grid_tergenang'], 0)} ({id_num(s['persen_grid_tergenang'], 2)} %); "
                 f"ruas ditutup {id_num(s['n_roads_closed'], 0)}; TES valid {id_num(sum(s['tes_valid'].values()), 0)}; "
                 f"rerata waktu minimum {id_num(s['mean_waktu_min'], 2)} menit (median {id_num(s['median_waktu_min'], 2)}); "
                 f"TAS {id_num(t['jumlah_tas'], 0)} ({id_num(t['persen_tas_non_tergenang'], 2)} % non-Tergenang); "
                 f"Terputus {id_num(t['jumlah_terputus'], 0)}; TAS persentil {id_num(t['sensitivitas_persentil']['jumlah_tas'], 0)}")
    L += ["", "## Pemilihan K", ""]
    for c in ks["candidates"]:
        if c["k"] in order:
            L.append(f"- K = {c['k']}: ARI subsampel {id_num(c['ari_subsampel_mean'], 3)} ± {id_num(c['ari_subsampel_sd'], 3)}; "
                     f"ARI inisialisasi {id_num(c['ari_inisialisasi_mean'], 3)}; Silhouette {id_num(c['silhouette'], 3)}")
    L.append(f"- Aturan stabilitas (pemecah seri K terkecil) memilih K = {ks['k_terpilih']}; kandidat dalam toleransi: "
             f"{', '.join(map(str, ks['k_dalam_toleransi']))}")
    L += ["", "## Model", ""]
    for k in order:
        L += _summary_k(R, k, grid)
    (out_dir / "ringkasan_angka_bab4.md").write_text("\n".join(L) + "\n", encoding="utf-8")


def write_v2_v3(R, grid, out_dir):
    if not ARSIP_V2.exists():
        return
    V2 = json.loads(ARSIP_V2.read_text(encoding="utf-8"))
    l2, l3 = V2["levels"], R["levels"]
    rows = []

    def add(nama, a, b, dec, sebab):
        rows.append(f"| {nama} | {id_num(a, dec)} | {id_num(b, dec)} | {sebab} |")

    SR = "Kelas bahaya dari raster InaRisk (grid mayoritas piksel, ruas maksimum)"
    for key in LEVEL_ORDER:
        add(f"Grid Tergenang — {LEVEL_LABEL[key]}", l2[key]["n_grid_tergenang"], l3[key]["n_grid_tergenang"], 0, SR)
    for key in LEVEL_ORDER:
        add(f"Ruas ditutup — {LEVEL_LABEL[key]}", l2[key]["n_roads_closed"], l3[key]["n_roads_closed"], 0,
            SR + "; aturan maksimum menaikkan kelas 2.477 ruas")
    for key in LEVEL_ORDER:
        add(f"TES valid — {LEVEL_LABEL[key]}", l2[key]["diagnostik"]["tes_valid_total"],
            l3[key]["diagnostik"]["tes_valid_total"], 0, "Kelas TES sama; aturan radius 50 m memakai grid Tergenang baru")
    add("T_pen (menit)", V2["t_pen"], R["t_pen"], 2, "Tidak ada perubahan pada Baseline (tidak ada penutupan)")
    for key in LEVEL_ORDER:
        add(f"Rerata waktu minimum non-Tergenang — {LEVEL_LABEL[key]}", l2[key]["mean_waktu_min"], l3[key]["mean_waktu_min"], 2,
            SR)
    add("Baris data gabungan", V2["data_gabungan"]["n_baris"], R["data_gabungan"]["n_baris"], 0, SR)
    add("Komponen PCA", V2["preprocessing"]["n_components"], R["preprocessing"]["n_components"], 0, "Data gabungan berubah")
    add("K menurut aturan stabilitas", V2["k"], R["k_selection"]["k_terpilih"], 0,
        "Stabilitas K dihitung ulang pada data v3 (keputusan peneliti)")
    c2 = {c["k"]: c for c in V2["k_selection"]["candidates"]}
    c3 = {c["k"]: c for c in R["k_selection"]["candidates"]}
    for k in (2, 3):
        add(f"ARI subsampel K = {k}", c2[k]["ari_subsampel_mean"], c3[k]["ari_subsampel_mean"], 3, "Data v3")
    for key in LEVEL_ORDER:
        t2, t3 = l2[key]["tas"], l3[key]["tas"]
        add(f"TAS — {LEVEL_LABEL[key]}", t2["jumlah_tas"], t3["jumlah_tas"], 0,
            "Aturan utama v3: DI ≥ 2 dan T_ideal ≤ 5 menit (v2: persentil P25/P75) + " + SR)
        add(f"TAS aturan persentil — {LEVEL_LABEL[key]}", t2["jumlah_tas"], t3["sensitivitas_persentil"]["jumlah_tas"], 0,
            "Aturan sama (persentil); perbedaan karena " + SR)
        add(f"Terputus — {LEVEL_LABEL[key]}", t2["jumlah_terputus"], t3["jumlah_terputus"], 0, SR)
    tr2 = {(t["from_level"], t["to_level"]): t for t in V2["transitions"]}
    for t in R["model"]["2"]["transitions"]:
        a = tr2[(t["from_level"], t["to_level"])]
        nm = f"{LEVEL_LABEL[t['from_level']]} → {LEVEL_LABEL[t['to_level']]}"
        add(f"Stability rate K = 2 {nm} (%)", a["stability_rate"], t["stability_rate"], 2, SR)
        add(f"ARI K = 2 {nm}", a["ari"], t["ari"], 3, SR)
    a2 = {a["algoritma"]: a for a in V2.get("algorithm_comparison", [])}
    for a in R["model"]["2"].get("algorithm_comparison", []):
        b = a2.get(a["algoritma"])
        sebab = "m SFCM 2,0 → 1,7 + " + SR if a["algoritma"] == "SFCM" else SR
        add(f"Silhouette {a['algoritma']} (K = 2)", b.get("silhouette") if b else None, a.get("silhouette"), 3, sebab)
    L = ["# Perbandingan angka kunci v2 vs v3", "",
         "v2 = `data/locked/arsip_v2/` (tag `hasil-skripsi-v2`); v3 = `data/locked/` (tag `hasil-skripsi-v3`). "
         "**Angka v3 yang berlaku.** Kolom terakhir = perubahan metode penyebab perbedaan.", "",
         "| Angka | v2 | v3 | Perubahan metode penyebab |", "|---|---:|---:|---|", *rows, "",
         "Perubahan metode v3 lainnya (tidak mengubah angka di atas secara langsung): label interpretasi berbasis median "
         "dan proporsi penalti; PC/PE; tabel K utama/lampiran; keluaran lengkap K = 2 dan K = 3."]
    (out_dir / "perbandingan_v2_vs_v3.md").write_text("\n".join(L) + "\n", encoding="utf-8")


def main(argv=None):
    ap = argparse.ArgumentParser(description="Ekspor tabel Bab IV dari data/locked/")
    ap.add_argument("--k-utama", type=int, choices=(2, 3), default=None)
    ap.add_argument("--keluaran", type=Path, default=None)
    args = ap.parse_args(argv)
    R, lock, grid = load_locked()
    k_utama = args.k_utama or k_utama_default()
    ks = [int(k) for k in R["k_keluaran"]]
    if k_utama not in ks:
        raise SystemExit(f"K utama {k_utama} tidak ada dalam keluaran {ks}")
    out_dir = args.keluaran or OUT
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in ("perbandingan_v1_vs_v2.md",):                 # laporan versi lama tersimpan di arsip_v2/
        (out_dir / old).unlink(missing_ok=True)
    tables = tables_common(R) + tables_for_k(R, k_utama, "T") + tables_k2_vs_k3(R, grid)
    for k in ks:
        if k != k_utama:
            tables += tables_for_k(R, k, "S")
    order = lambda t: ({"T": 0, "L": 1, "S": 2}[t.code[0]], int("".join(ch for ch in t.code[1:3] if ch.isdigit())), t.code)
    tables.sort(key=order)
    write_csv(tables, out_dir)
    write_docx(tables, lock, k_utama, out_dir)
    write_summary(R, lock, k_utama, grid, out_dir)
    write_v2_v3(R, grid, out_dir)
    write_maps(R, grid, out_dir)
    print(f"{len(tables)} tabel diekspor ke {out_dir} (K utama = {k_utama})")


if __name__ == "__main__":
    sys.exit(main())
